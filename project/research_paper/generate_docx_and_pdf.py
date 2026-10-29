"""
Academic Paper Formatter: Converts Markdown to Formatted DOCX and PDF
Outputs:
- project/research_paper/research_paper.docx
- project/research_paper/research_paper.pdf
"""

import re
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

PAPER_DIR = Path(__file__).resolve().parent
MD_PATH = PAPER_DIR / "research_paper.md"
DOCX_PATH = PAPER_DIR / "research_paper.docx"
PDF_PATH = PAPER_DIR / "research_paper.pdf"
DIAGRAM_PATH = PAPER_DIR / "methodology_diagram.png"

# ==============================================================================
# 1. GENERATE DOCX
# ==============================================================================
def generate_docx():
    print("Generating DOCX...")
    doc = Document()

    # Set 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()

    def add_table_to_doc(t_lines):
        rows = []
        for l in t_lines:
            l = l.strip()
            if not l or l.startswith("+-") or l.startswith("| :"):
                continue
            if l.startswith("|") and l.endswith("|"):
                cells = [c.strip() for c in l[1:-1].split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    continue
                rows.append(cells)
        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row in enumerate(rows):
            for c_idx, cell_text in enumerate(row):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    if r_idx == 0:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                        shading = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls("w")))
                        cell._tc.get_or_add_tcPr().append(shading)
                    else:
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        doc.add_paragraph()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Title
        if line.startswith("# ") and i == 0:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:])
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        # Author block
        if line.startswith("**Author Names Blinded"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("Author Names Blinded for Peer Review")
            run.font.size = Pt(11)
            run.font.bold = True
            i += 1
            continue

        if line.startswith("*Department") or line.startswith("*Affiliations") or line.startswith("Email:"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line.strip("*"))
            run.font.size = Pt(9.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            i += 1
            continue

        if line.startswith("---"):
            i += 1
            continue

        # ASCII diagram replacement with Figure 1 Image
        if line.startswith("+--") and ("1. MULTI-MODAL" in lines[min(i+1, len(lines)-1)] or "MULTICOMPONENT" in lines[min(i+1, len(lines)-1)]):
            while i < len(lines) and not lines[i].startswith("**Fig. 1"):
                i += 1
            if i < len(lines) and lines[i].startswith("**Fig. 1"):
                caption = lines[i].strip().replace("**", "")
                if DIAGRAM_PATH.exists():
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(12)
                    p_img.paragraph_format.space_after = Pt(4)
                    r_img = p_img.add_run()
                    r_img.add_picture(str(DIAGRAM_PATH), width=Inches(6.2))
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(12)
                    r_cap = p_cap.add_run(caption)
                    r_cap.font.size = Pt(9.5)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
                i += 1
            continue

        # Section Headings
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:])
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:])
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:])
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Tables
        if line.startswith("+--") or (line.startswith("|") and "|" in line[1:]):
            t_lines = [line]
            i += 1
            while i < len(lines) and (lines[i].startswith("+--") or lines[i].startswith("|")):
                t_lines.append(lines[i])
                i += 1
            add_table_to_doc(t_lines)
            continue

        # Lists
        if line.strip().startswith("- ") or re.match(r"^[0-9]+\.\s+", line.strip()):
            p = doc.add_paragraph(style="List Bullet" if line.strip().startswith("- ") else "List Number")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            clean = re.sub(r"^[0-9]+\.\s+|^-\s+", "", line.strip())
            parts = re.split(r"(\*\*.*?\*\*)", clean)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Regular paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith("*") and part.endswith("*"):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                else:
                    p.add_run(part)

        i += 1

    doc.save(str(DOCX_PATH))
    print(f"DOCX created successfully at: {DOCX_PATH}")


# ==============================================================================
# 2. GENERATE PDF VIA REPORTLAB
# ==============================================================================
def generate_pdf():
    print("Generating PDF...")
    target_pdf = PDF_PATH
    try:
        with open(target_pdf, "ab"):
            pass
    except (PermissionError, OSError):
        target_pdf = PAPER_DIR / "research_paper_v2.pdf"
        print(f"Notice: research_paper.pdf is open in another viewer. Writing to {target_pdf.name}")

    doc = SimpleDocTemplate(
        str(target_pdf),
        pagesize=letter,
        leftMargin=45,
        rightMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0f172a"),
        alignment=1, # Center
        spaceAfter=8
    )

    author_style = ParagraphStyle(
        "DocAuthor",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#334155"),
        alignment=1,
        spaceAfter=2
    )

    affil_style = ParagraphStyle(
        "DocAffil",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#64748b"),
        alignment=1,
        spaceAfter=10
    )

    h1_style = ParagraphStyle(
        "DocH1",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13.5,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=10,
        spaceAfter=3,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        "DocH2",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=7,
        spaceAfter=2,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=9.0,
        leading=11.5,
        textColor=colors.HexColor("#1a1a1a"),
        alignment=4, # Justify
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=9.0,
        leading=11.5,
        textColor=colors.HexColor("#1a1a1a"),
        leftIndent=14,
        spaceAfter=2
    )

    caption_style = ParagraphStyle(
        "DocCaption",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8.0,
        leading=10.5,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceBefore=3,
        spaceAfter=8
    )

    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=7.8,
        leading=9.2,
        textColor=colors.HexColor("#1e293b")
    )

    table_hdr_style = ParagraphStyle(
        "TableHdr",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.8,
        leading=9.2,
        textColor=colors.HexColor("#0f172a")
    )

    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()
    story = []

    def make_rl_table(t_lines):
        raw_rows = []
        for l in t_lines:
            l = l.strip()
            if not l or l.startswith("+-") or l.startswith("| :"):
                continue
            if l.startswith("|") and l.endswith("|"):
                cells = [c.strip() for c in l[1:-1].split("|")]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    continue
                raw_rows.append(cells)
        if not raw_rows:
            return None

        col_count = max(len(r) for r in raw_rows)
        # 522 pt available width (612 - 2*45)
        col_width = 522.0 / col_count

        table_data = []
        for r_idx, r in enumerate(raw_rows):
            row_data = []
            for c_idx in range(col_count):
                txt = r[c_idx] if c_idx < len(r) else ""
                txt_clean = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", txt)
                txt_clean = re.sub(r"\*(.*?)\*", r"<i>\1</i>", txt_clean)
                style_to_use = table_hdr_style if r_idx == 0 else table_cell_style
                row_data.append(Paragraph(txt_clean, style_to_use))
            table_data.append(row_data)

        t = Table(table_data, colWidths=[col_width]*col_count)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#F1F5F9")),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
            ('LINEBELOW', (0,0), (-1,0), 1.2, colors.HexColor("#0F172A")),
        ]))
        return t

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Title
        if line.startswith("# ") and i == 0:
            story.append(Paragraph(line[2:], title_style))
            i += 1
            continue

        if line.startswith("**Author Names Blinded"):
            story.append(Paragraph("Author Names Blinded for Peer Review", author_style))
            i += 1
            continue

        if line.startswith("*Department") or line.startswith("*Affiliations") or line.startswith("Email:"):
            story.append(Paragraph(line.strip("*"), affil_style))
            i += 1
            continue

        if line.startswith("---"):
            i += 1
            continue

        # Diagram
        if line.startswith("+--") and ("1. MULTI-MODAL" in lines[min(i+1, len(lines)-1)] or "MULTICOMPONENT" in lines[min(i+1, len(lines)-1)]):
            while i < len(lines) and not lines[i].startswith("**Fig. 1"):
                i += 1
            if i < len(lines) and lines[i].startswith("**Fig. 1"):
                caption = lines[i].strip().replace("**", "")
                if DIAGRAM_PATH.exists():
                    story.append(Spacer(1, 4))
                    img = Image(str(DIAGRAM_PATH), width=522, height=240)
                    story.append(img)
                    story.append(Paragraph(caption, caption_style))
                i += 1
            continue

        # Headings
        if line.startswith("## "):
            story.append(Paragraph(line[3:], h1_style))
            i += 1
            continue

        if line.startswith("### "):
            story.append(Paragraph(line[4:], h2_style))
            i += 1
            continue

        if line.startswith("#### "):
            story.append(Paragraph(f"<i><b>{line[5:]}</b></i>", h2_style))
            i += 1
            continue

        # Tables
        if line.startswith("+--") or (line.startswith("|") and "|" in line[1:]):
            t_lines = [line]
            i += 1
            while i < len(lines) and (lines[i].startswith("+--") or lines[i].startswith("|")):
                t_lines.append(lines[i])
                i += 1
            rl_t = make_rl_table(t_lines)
            if rl_t:
                story.append(Spacer(1, 4))
                story.append(rl_t)
                story.append(Spacer(1, 6))
            continue

        # Lists
        if line.strip().startswith("- ") or re.match(r"^[0-9]+\.\s+", line.strip()):
            clean = line.strip()
            clean = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", clean)
            clean = re.sub(r"\*(.*?)\*", r"<i>\1</i>", clean)
            story.append(Paragraph(clean, bullet_style))
            i += 1
            continue

        # Body paragraph
        if line.strip():
            clean = line.strip()
            clean = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", clean)
            clean = re.sub(r"\*(.*?)\*", r"<i>\1</i>", clean)
            story.append(Paragraph(clean, body_style))

        i += 1

    doc.build(story)
    print(f"PDF created successfully at: {target_pdf}")


if __name__ == "__main__":
    generate_docx()
    generate_pdf()
